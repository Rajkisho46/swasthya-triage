import os
import pptx
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

SCREENSHOTS_DIR = os.path.join(os.path.dirname(__file__), "presentation_screenshots")
OUTPUT_PPTX = os.path.join(os.path.dirname(__file__), "Swasthya_Triage_Hackathon_Presentation.pptx")

# Color Palette: AI Healthcare Command Center
BG_DARK = RGBColor(7, 13, 30)         # #070d1e Deep Midnight Blue
CARD_BG = RGBColor(15, 27, 53)        # #0f1b35 Navy Surface
CARD_BORDER = RGBColor(30, 48, 86)    # #1e3056 Card Border
TEAL_ACCENT = RGBColor(0, 229, 255)   # #00e5ff Bright Cyan / Teal
MINT_ACCENT = RGBColor(52, 211, 153)  # #34d399 Soft Mint
CORAL_ACCENT = RGBColor(248, 113, 113)# #f87171 Red-Flag Alert
BLUE_GLOW = RGBColor(56, 189, 248)    # #38bdf8 Sky Blue
TEXT_WHITE = RGBColor(255, 255, 255)  # #ffffff Primary Text
TEXT_LIGHT = RGBColor(203, 213, 225)  # #cbd5e1 Light Slate
TEXT_MUTED = RGBColor(148, 163, 184)  # #94a3b8 Muted Text
BADGE_BG = RGBColor(12, 38, 70)       # #0c2646 Pill Badge Navy
ACCENT_CARD = RGBColor(18, 35, 68)

def create_presentation():
    prs = Presentation()
    # 16:9 Widescreen dimensions: 13.333 x 7.5 inches
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    def set_slide_background(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
        bg.fill.solid()
        bg.fill.fore_color.rgb = BG_DARK
        bg.line.fill.background()
        return bg

    def add_header(slide, title, category="SWASTHYA TRIAGE  •  HACKATHON DECK", slide_num=None):
        # Category Pill
        pill = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.4), Inches(4.5), Inches(0.32))
        pill.fill.solid()
        pill.fill.fore_color.rgb = BADGE_BG
        pill.line.color.rgb = BLUE_GLOW
        pill.line.width = Pt(1)
        tf_pill = pill.text_frame
        tf_pill.word_wrap = True
        p_pill = tf_pill.paragraphs[0]
        p_pill.text = f"  {category.upper()}"
        p_pill.font.size = Pt(9.5)
        p_pill.font.bold = True
        p_pill.font.color.rgb = TEAL_ACCENT
        p_pill.font.name = 'Segoe UI'

        # Main Title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.72), Inches(11.7), Inches(0.65))
        tf = title_box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(22)
        p.font.bold = True
        p.font.color.rgb = TEXT_WHITE
        p.font.name = 'Segoe UI'

        # Slide Number
        if slide_num:
            num_box = slide.shapes.add_textbox(Inches(11.8), Inches(0.4), Inches(0.9), Inches(0.4))
            p_num = num_box.text_frame.paragraphs[0]
            p_num.text = f"{slide_num:02d} / 10"
            p_num.font.size = Pt(11)
            p_num.font.bold = True
            p_num.font.color.rgb = TEXT_MUTED
            p_num.font.name = 'Segoe UI'
            p_num.alignment = PP_ALIGN.RIGHT

    # =========================================================================
    # SLIDE 1 — TITLE / PROBLEM STATEMENT
    # =========================================================================
    s1 = prs.slides.add_slide(blank_layout)
    set_slide_background(s1)

    # Accent decorative glow bar
    bar = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.9), Inches(11.733), Inches(0.06))
    bar.fill.solid()
    bar.fill.fore_color.rgb = TEAL_ACCENT
    bar.line.fill.background()

    # Category badge
    cat1 = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.2), Inches(5.2), Inches(0.38))
    cat1.fill.solid()
    cat1.fill.fore_color.rgb = BADGE_BG
    cat1.line.color.rgb = BLUE_GLOW
    cat1.line.width = Pt(1)
    p_c1 = cat1.text_frame.paragraphs[0]
    p_c1.text = "  CLINICAL DECISION SUPPORT & MULTIMODAL TRIAGE"
    p_c1.font.size = Pt(10)
    p_c1.font.bold = True
    p_c1.font.color.rgb = TEAL_ACCENT
    p_c1.font.name = 'Segoe UI'

    # Project Title
    t1 = s1.shapes.add_textbox(Inches(0.8), Inches(1.65), Inches(8.5), Inches(1.2))
    p1 = t1.text_frame.paragraphs[0]
    p1.text = "SWASTHYA TRIAGE"
    p1.font.size = Pt(40)
    p1.font.bold = True
    p1.font.color.rgb = TEXT_WHITE
    p1.font.name = 'Segoe UI'

    sub1 = t1.text_frame.add_paragraph()
    sub1.text = "Clinician-Supervised Multimodal Patient Intake & Safe Triage Assistance"
    sub1.font.size = Pt(15)
    sub1.font.color.rgb = BLUE_GLOW
    sub1.font.name = 'Segoe UI'

    # Left Card: Hackathon Entry Fields
    fcard = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(3.1), Inches(6.0), Inches(3.8))
    fcard.fill.solid()
    fcard.fill.fore_color.rgb = CARD_BG
    fcard.line.color.rgb = CARD_BORDER
    fcard.line.width = Pt(1)

    tf_f = fcard.text_frame
    tf_f.word_wrap = True
    fields = [
        ("Problem Statement:", "[ADD LATER]"),
        ("Problem Statement ID:", "[ADD LATER]"),
        ("Team Name:", "[ADD LATER]"),
        ("Domain:", "[ADD LATER]"),
        ("Tagline:", "[ADD LATER]"),
        ("Architecture:", "FastAPI + React 19 + AI Safety Pipeline (Non-Diagnostic)")
    ]
    for idx, (lbl, val) in enumerate(fields):
        p_l = tf_f.paragraphs[0] if idx == 0 else tf_f.add_paragraph()
        p_l.text = f"{lbl} "
        p_l.font.size = Pt(11)
        p_l.font.bold = True
        p_l.font.color.rgb = TEAL_ACCENT if idx < 5 else MINT_ACCENT
        p_l.font.name = 'Segoe UI'
        
        run_v = p_l.add_run()
        run_v.text = val
        run_v.font.size = Pt(11)
        run_v.font.bold = False
        run_v.font.color.rgb = TEXT_WHITE if idx == 5 else RGBColor(253, 224, 71)
        run_v.font.name = 'Segoe UI'
        p_l.space_after = Pt(8)

    # Right Card: Real Screenshot
    sc1 = os.path.join(SCREENSHOTS_DIR, "01_patient_login.png")
    if os.path.exists(sc1):
        # Frame
        frame1 = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.1), Inches(3.1), Inches(5.433), Inches(3.8))
        frame1.fill.solid()
        frame1.fill.fore_color.rgb = CARD_BG
        frame1.line.color.rgb = BLUE_GLOW
        frame1.line.width = Pt(1.5)
        # Image inside
        s1.shapes.add_picture(sc1, Inches(7.2), Inches(3.2), Inches(5.233), Inches(3.6))

    # =========================================================================
    # SLIDE 2 — PROBLEM & MOTIVATION
    # =========================================================================
    s2 = prs.slides.add_slide(blank_layout)
    set_slide_background(s2)
    add_header(s2, "The Problem: Information Overload & Delayed Urgency Identification", "PROBLEM & MOTIVATION", 2)

    # Left Column: Key Pain Points (3 Cards)
    pain_points = [
        ("Unstructured Patient Information", "Patients express distress in colloquial, disorganized narratives across text, speech, or paper reports where critical red-flags are easily obscured."),
        ("Manual & Delayed Prioritization", "Emergency and outpatient queues struggle to rapidly distinguish routine complaints from time-sensitive emergencies before bedside assessment."),
        ("Lack of Clinician-Supervised AI", "Conventional AI chat systems pose high hallucination risk. Healthcare needs strictly non-diagnostic AI with human clinician override.")
    ]
    y_card = 1.45
    for title_p, desc_p in pain_points:
        p_card = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(y_card), Inches(5.7), Inches(1.3))
        p_card.fill.solid()
        p_card.fill.fore_color.rgb = CARD_BG
        p_card.line.color.rgb = CARD_BORDER
        p_card.line.width = Pt(1)
        
        tf = p_card.text_frame
        tf.word_wrap = True
        p_t = tf.paragraphs[0]
        p_t.text = f"• {title_p}"
        p_t.font.size = Pt(12)
        p_t.font.bold = True
        p_t.font.color.rgb = TEAL_ACCENT
        p_t.font.name = 'Segoe UI'
        
        p_d = tf.add_paragraph()
        p_d.text = desc_p
        p_d.font.size = Pt(10.5)
        p_d.font.color.rgb = TEXT_LIGHT
        p_d.font.name = 'Segoe UI'
        p_d.space_before = Pt(3)
        y_card += 1.42

    # Right Column: Flow Comparison + Real Screenshot
    # Current vs Our Approach Box
    f_comp = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.45), Inches(5.733), Inches(2.2))
    f_comp.fill.solid()
    f_comp.fill.fore_color.rgb = ACCENT_CARD
    f_comp.line.color.rgb = CARD_BORDER
    f_comp.line.width = Pt(1)
    
    tf_c = f_comp.text_frame
    tf_c.word_wrap = True
    
    p_cc = tf_c.paragraphs[0]
    p_cc.text = "CURRENT CHALLENGE (Traditional Queue):"
    p_cc.font.size = Pt(10.5)
    p_cc.font.bold = True
    p_cc.font.color.rgb = CORAL_ACCENT
    p_cc.font.name = 'Segoe UI'
    
    p_cc_flow = tf_c.add_paragraph()
    p_cc_flow.text = "Patient  →  Unstructured Info  →  Manual Interpretation  →  Delayed Prioritization"
    p_cc_flow.font.size = Pt(10)
    p_cc_flow.font.color.rgb = TEXT_MUTED
    p_cc_flow.font.name = 'Segoe UI'
    p_cc_flow.space_after = Pt(10)
    
    p_oa = tf_c.add_paragraph()
    p_oa.text = "OUR APPROACH (Swasthya Triage):"
    p_oa.font.size = Pt(10.5)
    p_oa.font.bold = True
    p_oa.font.color.rgb = MINT_ACCENT
    p_oa.font.name = 'Segoe UI'
    
    p_oa_flow = tf_c.add_paragraph()
    p_oa_flow.text = "Patient  →  Structured Multimodal Intake  →  AI Safety & Triage  →  Human Clinical Review"
    p_oa_flow.font.size = Pt(10)
    p_oa_flow.font.bold = True
    p_oa_flow.font.color.rgb = TEXT_WHITE
    p_oa_flow.font.name = 'Segoe UI'

    # Screenshot below
    sc2 = os.path.join(SCREENSHOTS_DIR, "02_patient_portal_dashboard.png")
    if os.path.exists(sc2):
        s2.shapes.add_picture(sc2, Inches(6.8), Inches(3.8), Inches(5.733), Inches(3.1))

    # =========================================================================
    # SLIDE 3 — OUR SOLUTION
    # =========================================================================
    s3 = prs.slides.add_slide(blank_layout)
    set_slide_background(s3)
    add_header(s3, "Introducing Swasthya Triage: Structured Multimodal Intake & Clinical Decision Support", "OUR SOLUTION", 3)

    # Core Value Banner
    banner = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.4), Inches(11.733), Inches(0.5))
    banner.fill.solid()
    banner.fill.fore_color.rgb = BADGE_BG
    banner.line.color.rgb = TEAL_ACCENT
    banner.line.width = Pt(1)
    p_b = banner.text_frame.paragraphs[0]
    p_b.text = "   CORE PRINCIPLE:  \"AI assists the workflow. Clinicians make the final decision.\""
    p_b.font.size = Pt(11.5)
    p_b.font.bold = True
    p_b.font.color.rgb = TEAL_ACCENT
    p_b.font.name = 'Segoe UI'

    # Horizontal 6-Stage Workflow Pipeline
    stages = [
        ("01. Patient Intake", "Multimodal input via Text, Voice, or Lab OCR"),
        ("02. Health AI", "Symptom extraction & deterministic red-flag check"),
        ("03. Priority Triage", "Urgency signal scoring & clinical evidence card"),
        ("04. Nurse Queue", "Prioritized triaging & bedside vitals recording"),
        ("05. Reviewer MD", "Clinical confirmation, modification, or override"),
        ("06. Referral & Audit", "Instant PDF generation & immutable event log")
    ]
    w_box = 1.83
    x_box = 0.8
    for idx, (st_t, st_d) in enumerate(stages):
        box_s = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x_box), Inches(2.05), Inches(w_box), Inches(1.4))
        box_s.fill.solid()
        box_s.fill.fore_color.rgb = CARD_BG
        box_s.line.color.rgb = TEAL_ACCENT if idx in [1, 2] else BLUE_GLOW
        box_s.line.width = Pt(1)
        
        tf_s = box_s.text_frame
        tf_s.word_wrap = True
        p_st = tf_s.paragraphs[0]
        p_st.text = st_t
        p_st.font.size = Pt(10)
        p_st.font.bold = True
        p_st.font.color.rgb = TEAL_ACCENT if idx in [1, 2] else TEXT_WHITE
        p_st.font.name = 'Segoe UI'
        
        p_sd = tf_s.add_paragraph()
        p_sd.text = st_d
        p_sd.font.size = Pt(8.5)
        p_sd.font.color.rgb = TEXT_LIGHT
        p_sd.font.name = 'Segoe UI'
        p_sd.space_before = Pt(3)
        
        x_box += 1.98

    # Bottom Grid: Features List + 2 Real Screenshots
    f_list_card = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(3.6), Inches(4.5), Inches(3.4))
    f_list_card.fill.solid()
    f_list_card.fill.fore_color.rgb = CARD_BG
    f_list_card.line.color.rgb = CARD_BORDER
    f_list_card.line.width = Pt(1)
    
    tf_fl = f_list_card.text_frame
    tf_fl.word_wrap = True
    p_flh = tf_fl.paragraphs[0]
    p_flh.text = "Key Implemented Capabilities:"
    p_flh.font.size = Pt(12)
    p_flh.font.bold = True
    p_flh.font.color.rgb = TEAL_ACCENT
    p_flh.font.name = 'Segoe UI'
    
    caps = [
        "Patient Self-Intake & Authentication (JWT / BCrypt)",
        "Conversational Health AI with Safety Guardrails",
        "Deterministic Urgency Rules (Cardiorespiratory, etc.)",
        "Role-Based Access Control (Patient, Nurse, Doctor, Admin)",
        "Nurse Bedside Vitals (Temp, BP, SpO2, Heart Rate)",
        "Doctor Review & Direct Override Protocol",
        "Referral Letter Generation with ReportLab PDF"
    ]
    for c in caps:
        p_c = tf_fl.add_paragraph()
        p_c.text = f"✔  {c}"
        p_c.font.size = Pt(9.5)
        p_c.font.color.rgb = TEXT_LIGHT
        p_c.font.name = 'Segoe UI'
        p_c.space_before = Pt(2)

    # 2 Screenshots
    sc3a = os.path.join(SCREENSHOTS_DIR, "04_patient_intake_multimodal.png")
    if os.path.exists(sc3a):
        s3.shapes.add_picture(sc3a, Inches(5.45), Inches(3.6), Inches(3.4), Inches(3.4))

    sc3b = os.path.join(SCREENSHOTS_DIR, "08_medical_reviewer_workspace.png")
    if os.path.exists(sc3b):
        s3.shapes.add_picture(sc3b, Inches(9.0), Inches(3.6), Inches(3.533), Inches(3.4))

    # =========================================================================
    # SLIDE 4 — HOW THE HEALTH AI WORKS
    # =========================================================================
    s4 = prs.slides.add_slide(blank_layout)
    set_slide_background(s4)
    add_header(s4, "Health AI — From Patient Input to Safe Assistance", "AI & SAFETY ENGINE", 4)

    # Left Column: AI Pipeline Flowchart (Step-by-step boxes)
    ai_steps = [
        ("1. Multimodal Intake", "Patient submits text complaint, voice narration, or lab report photo."),
        ("2. Clinical Structuring", "Extracts Chief Complaint, Symptoms, Timeline, and Medical History."),
        ("3. Safety & Red-Flag Rules", "Evaluates deterministic emergency rules (e.g. Chest pain + dyspnea)."),
        ("4. Triage Prioritization", "Computes Urgency Level (Emergency, Urgent, Routine) with evidence."),
        ("5. Guardrailed Output", "Generates non-diagnostic guidance with mandatory clinician review warning.")
    ]
    y_step = 1.45
    for s_title, s_desc in ai_steps:
        s_box = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(y_step), Inches(5.8), Inches(1.0))
        s_box.fill.solid()
        s_box.fill.fore_color.rgb = CARD_BG
        s_box.line.color.rgb = BLUE_GLOW if "Safety" in s_title else CARD_BORDER
        s_box.line.width = Pt(1)
        
        tf_s = s_box.text_frame
        tf_s.word_wrap = True
        p_st = tf_s.paragraphs[0]
        p_st.text = s_title
        p_st.font.size = Pt(11)
        p_st.font.bold = True
        p_st.font.color.rgb = TEAL_ACCENT
        p_st.font.name = 'Segoe UI'
        
        p_sd = tf_s.add_paragraph()
        p_sd.text = s_desc
        p_sd.font.size = Pt(9.5)
        p_sd.font.color.rgb = TEXT_LIGHT
        p_sd.font.name = 'Segoe UI'
        y_step += 1.08

    # Right Column: Health AI Chatbox Screenshot + Safety Disclaimer Card
    sc4 = os.path.join(SCREENSHOTS_DIR, "03_health_ai_chat.png")
    if os.path.exists(sc4):
        s4.shapes.add_picture(sc4, Inches(6.9), Inches(1.45), Inches(5.633), Inches(4.3))

    safe_box = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.9), Inches(5.9), Inches(5.633), Inches(1.15))
    safe_box.fill.solid()
    safe_box.fill.fore_color.rgb = BADGE_BG
    safe_box.line.color.rgb = CORAL_ACCENT
    safe_box.line.width = Pt(1)
    
    tf_sb = safe_box.text_frame
    tf_sb.word_wrap = True
    p_sbh = tf_sb.paragraphs[0]
    p_sbh.text = "🛡️ NON-DIAGNOSTIC SAFETY GUARANTEE:"
    p_sbh.font.size = Pt(10)
    p_sbh.font.bold = True
    p_sbh.font.color.rgb = CORAL_ACCENT
    p_sbh.font.name = 'Segoe UI'
    
    p_sbd = tf_sb.add_paragraph()
    p_sbd.text = "The AI assists with structuring information and identifying urgency signals. It DOES NOT independently diagnose medical conditions or prescribe medications."
    p_sbd.font.size = Pt(9)
    p_sbd.font.color.rgb = TEXT_LIGHT
    p_sbd.font.name = 'Segoe UI'

    # =========================================================================
    # SLIDE 5 — MULTIMODAL HEALTHCARE INPUT
    # =========================================================================
    s5 = prs.slides.add_slide(blank_layout)
    set_slide_background(s5)
    add_header(s5, "Multimodal Patient Interaction: Unified Clinical Context", "MULTIMODAL INTAKE", 5)

    # Top: 4 Input Modes Grid
    modes = [
        ("1. TEXT INTAKE", "Natural language complaint in English/Hindi with duration and severity tracking.", BLUE_GLOW),
        ("2. VOICE ASSIST", "Speech-to-text audio recording for elderly and accessibility-first patients.", TEAL_ACCENT),
        ("3. MEDICAL OCR", "Lab report & prescription upload with OCR extraction of vitals/findings.", MINT_ACCENT),
        ("4. CLINICAL IMAGES", "Visual symptom attachment analyzed alongside patient clinical history.", RGBColor(168, 85, 247))
    ]
    w_m = 2.76
    x_m = 0.8
    for m_t, m_d, m_col in modes:
        m_box = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x_m), Inches(1.4), Inches(w_m), Inches(1.6))
        m_box.fill.solid()
        m_box.fill.fore_color.rgb = CARD_BG
        m_box.line.color.rgb = m_col
        m_box.line.width = Pt(1.5)
        
        tf_m = m_box.text_frame
        tf_m.word_wrap = True
        p_mt = tf_m.paragraphs[0]
        p_mt.text = m_t
        p_mt.font.size = Pt(11)
        p_mt.font.bold = True
        p_mt.font.color.rgb = m_col
        p_mt.font.name = 'Segoe UI'
        
        p_md = tf_m.add_paragraph()
        p_md.text = m_d
        p_md.font.size = Pt(9.5)
        p_md.font.color.rgb = TEXT_LIGHT
        p_md.font.name = 'Segoe UI'
        p_md.space_before = Pt(4)
        
        x_m += 2.98

    # Middle: Pipeline Indicator
    p_bar = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(3.15), Inches(11.733), Inches(0.45))
    p_bar.fill.solid()
    p_bar.fill.fore_color.rgb = BADGE_BG
    p_bar.line.color.rgb = BLUE_GLOW
    p_bar.line.width = Pt(1)
    p_pb = p_bar.text_frame.paragraphs[0]
    p_pb.text = "  [Text + Voice + Reports + Images]  ──▶  Input Processing Pipeline  ──▶  Single Unified Structured Clinical Context  "
    p_pb.font.size = Pt(10)
    p_pb.font.bold = True
    p_pb.font.color.rgb = TEAL_ACCENT
    p_pb.font.name = 'Segoe UI'
    p_pb.alignment = PP_ALIGN.CENTER

    # Bottom: 2 Screenshots
    sc5a = os.path.join(SCREENSHOTS_DIR, "04_patient_intake_multimodal.png")
    if os.path.exists(sc5a):
        s5.shapes.add_picture(sc5a, Inches(0.8), Inches(3.75), Inches(5.7), Inches(3.3))

    sc5b = os.path.join(SCREENSHOTS_DIR, "03_health_ai_chat.png")
    if os.path.exists(sc5b):
        s5.shapes.add_picture(sc5b, Inches(6.8), Inches(3.75), Inches(5.733), Inches(3.3))

    # =========================================================================
    # SLIDE 6 — TRIAGE & CLINICAL WORKFLOW
    # =========================================================================
    s6 = prs.slides.add_slide(blank_layout)
    set_slide_background(s6)
    add_header(s6, "From Patient Intake to Hospital Clinical Workflow", "CLINICAL WORKFLOW & RBAC", 6)

    # Left: Flowchart and Responsibility Matrix
    flow_steps = [
        ("PATIENT INTAKE", "Submits multimodal symptoms & history via secure portal.", BLUE_GLOW),
        ("AI-ASSISTED TRIAGE", "Extracts timeline, matches red-flag rules, computes urgency.", TEAL_ACCENT),
        ("NURSE QUEUE & BEDSIDE", "Nurse reviews prioritized list, takes bedside vitals (BP, SpO2).", MINT_ACCENT),
        ("MEDICAL REVIEWER (MD)", "Physician verifies AI triage, adds clinical notes, confirms decision.", RGBColor(245, 158, 11)),
        ("REFERRAL & CARE PATH", "Generates institutional referral PDF or schedules direct consult.", RGBColor(168, 85, 247))
    ]
    y_fl = 1.45
    for fl_t, fl_d, fl_c in flow_steps:
        box_f = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(y_fl), Inches(5.3), Inches(0.88))
        box_f.fill.solid()
        box_f.fill.fore_color.rgb = CARD_BG
        box_f.line.color.rgb = fl_c
        box_f.line.width = Pt(1)
        
        tf = box_f.text_frame
        tf.word_wrap = True
        p_t = tf.paragraphs[0]
        p_t.text = fl_t
        p_t.font.size = Pt(10)
        p_t.font.bold = True
        p_t.font.color.rgb = fl_c
        p_t.font.name = 'Segoe UI'
        
        p_d = tf.add_paragraph()
        p_d.text = fl_d
        p_d.font.size = Pt(8.5)
        p_d.font.color.rgb = TEXT_LIGHT
        p_d.font.name = 'Segoe UI'
        y_fl += 0.95

    # Responsibility Contrast Box at bottom left
    resp_box = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(6.25), Inches(5.3), Inches(0.85))
    resp_box.fill.solid()
    resp_box.fill.fore_color.rgb = BADGE_BG
    resp_box.line.color.rgb = CARD_BORDER
    resp_box.line.width = Pt(1)
    
    tf_rb = resp_box.text_frame
    tf_rb.word_wrap = True
    p_ai_r = tf_rb.paragraphs[0]
    p_ai_r.text = "AI ROLE:  Assists • Structures • Flags • Summarizes"
    p_ai_r.font.size = Pt(9.5)
    p_ai_r.font.bold = True
    p_ai_r.font.color.rgb = TEAL_ACCENT
    p_ai_r.font.name = 'Segoe UI'
    
    p_cl_r = tf_rb.add_paragraph()
    p_cl_r.text = "CLINICIAN ROLE:  Reviews • Overrides • Prescribes • Decides"
    p_cl_r.font.size = Pt(9.5)
    p_cl_r.font.bold = True
    p_cl_r.font.color.rgb = MINT_ACCENT
    p_cl_r.font.name = 'Segoe UI'

    # Right: 2 Real Screenshots (Nurse Queue + Nurse Bedside)
    sc6a = os.path.join(SCREENSHOTS_DIR, "06_nurse_triage_queue.png")
    if os.path.exists(sc6a):
        s6.shapes.add_picture(sc6a, Inches(6.35), Inches(1.45), Inches(6.183), Inches(2.75))

    sc6b = os.path.join(SCREENSHOTS_DIR, "07_nurse_bedside_assessment.png")
    if os.path.exists(sc6b):
        s6.shapes.add_picture(sc6b, Inches(6.35), Inches(4.35), Inches(6.183), Inches(2.75))

    # =========================================================================
    # SLIDE 7 — TECHNICAL ARCHITECTURE
    # =========================================================================
    s7 = prs.slides.add_slide(blank_layout)
    set_slide_background(s7)
    add_header(s7, "System Architecture: Robust Full-Stack Engineering", "TECHNICAL ARCHITECTURE", 7)

    # 4 Architectural Tier Cards (Horizontal or Grid)
    arch_tiers = [
        ("FRONTEND LAYER", "React 19 + TypeScript + Vite", [
            "Role-based portal routing (Patient, Nurse, Doctor, Admin)",
            "Glassmorphism UI system with dynamic scroll reveals",
            "Lucide icons, responsive tables & accessible forms"
        ], BLUE_GLOW),
        ("BACKEND API LAYER", "FastAPI + Python 3.12 (Asynchronous)", [
            "RESTful API with Pydantic validation & strict type schemas",
            "JWT Authentication (HS256) & BCrypt password hashing",
            "ReportLab PDF referral generation engine"
        ], TEAL_ACCENT),
        ("AI & SAFETY ENGINE", "Provider Abstraction + Deterministic Fallback", [
            "Dual-engine: LLM structured extraction + Rule engine",
            "Deterministic fallback when offline/unreachable",
            "Evidence card mapping & urgency signal generation"
        ], MINT_ACCENT),
        ("DATABASE & AUDIT", "SQLite (Local) / PostgreSQL Ready", [
            "Async SQLAlchemy ORM with full relational schema",
            "Persistent cases, consents, reviews, and bedside vitals",
            "Immutable audit event log for clinical provenance"
        ], RGBColor(168, 85, 247))
    ]
    w_t = 2.76
    x_t = 0.8
    for t_title, t_tech, t_points, t_col in arch_tiers:
        t_card = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x_t), Inches(1.45), Inches(w_t), Inches(4.2))
        t_card.fill.solid()
        t_card.fill.fore_color.rgb = CARD_BG
        t_card.line.color.rgb = t_col
        t_card.line.width = Pt(1.5)
        
        tf = t_card.text_frame
        tf.word_wrap = True
        p_h = tf.paragraphs[0]
        p_h.text = t_title
        p_h.font.size = Pt(11)
        p_h.font.bold = True
        p_h.font.color.rgb = t_col
        p_h.font.name = 'Segoe UI'
        
        p_sub = tf.add_paragraph()
        p_sub.text = t_tech
        p_sub.font.size = Pt(9.5)
        p_sub.font.bold = True
        p_sub.font.color.rgb = TEXT_WHITE
        p_sub.font.name = 'Segoe UI'
        p_sub.space_after = Pt(8)
        
        for pt in t_points:
            p_p = tf.add_paragraph()
            p_p.text = f"• {pt}"
            p_p.font.size = Pt(8.5)
            p_p.font.color.rgb = TEXT_LIGHT
            p_p.font.name = 'Segoe UI'
            p_p.space_before = Pt(3)
        
        x_t += 2.98

    # Bottom Full-Width Architecture Call Chain
    achain = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(5.8), Inches(11.733), Inches(1.2))
    achain.fill.solid()
    achain.fill.fore_color.rgb = BADGE_BG
    achain.line.color.rgb = CARD_BORDER
    achain.line.width = Pt(1)
    
    tf_ac = achain.text_frame
    tf_ac.word_wrap = True
    p_ach = tf_ac.paragraphs[0]
    p_ach.text = "COMPLETE SYSTEM DATA FLOW:"
    p_ach.font.size = Pt(10)
    p_ach.font.bold = True
    p_ach.font.color.rgb = TEAL_ACCENT
    p_ach.font.name = 'Segoe UI'
    
    p_ac_flow = tf_ac.add_paragraph()
    p_ac_flow.text = "Patient / Staff Browser  ──▶  Vite React UI  ──▶  FastAPI Endpoints (/api/intake, /api/triage, /api/review)  ──▶  AI Safety Pipeline  ──▶  SQLite / PostgreSQL  ──▶  Immutable Audit Log"
    p_ac_flow.font.size = Pt(9.5)
    p_ac_flow.font.bold = True
    p_ac_flow.font.color.rgb = TEXT_WHITE
    p_ac_flow.font.name = 'Segoe UI'
    p_ac_flow.space_before = Pt(4)

    # =========================================================================
    # SLIDE 8 — SAFETY, PRIVACY & HUMAN OVERSIGHT
    # =========================================================================
    s8 = prs.slides.add_slide(blank_layout)
    set_slide_background(s8)
    add_header(s8, "Safety by Design: Multi-Layer Clinical Oversight & Governance", "SAFETY & GOVERNANCE", 8)

    # Left 3 Safety Pillar Cards
    pillars = [
        ("1. Clinician Supervision & Override", "Doctors retain full clinical authority to confirm, modify, or completely override AI triage recommendations with mandatory clinical rationale notes."),
        ("2. Deterministic Red-Flag Rules", "Hardcoded medical safety rules immediately intercept acute cardiorespiratory and critical red flags, bypassing generic conversational delays."),
        ("3. Role-Based Access & Auditability", "Strict JWT-based RBAC separates Patient self-intake from Nurse and Doctor workflows. All actions are written to an immutable audit trail.")
    ]
    y_pil = 1.45
    for pil_t, pil_d in pillars:
        pil_card = s8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(y_pil), Inches(5.7), Inches(1.3))
        pil_card.fill.solid()
        pil_card.fill.fore_color.rgb = CARD_BG
        pil_card.line.color.rgb = TEAL_ACCENT if "Override" in pil_t else CARD_BORDER
        pil_card.line.width = Pt(1)
        
        tf = pil_card.text_frame
        tf.word_wrap = True
        p_t = tf.paragraphs[0]
        p_t.text = pil_t
        p_t.font.size = Pt(11)
        p_t.font.bold = True
        p_t.font.color.rgb = TEAL_ACCENT
        p_t.font.name = 'Segoe UI'
        
        p_d = tf.add_paragraph()
        p_d.text = pil_d
        p_d.font.size = Pt(9.5)
        p_d.font.color.rgb = TEXT_LIGHT
        p_d.font.name = 'Segoe UI'
        p_d.space_before = Pt(3)
        y_pil += 1.42

    # Bottom statement
    gov_card = s8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(5.8), Inches(5.7), Inches(1.2))
    gov_card.fill.solid()
    gov_card.fill.fore_color.rgb = BADGE_BG
    gov_card.line.color.rgb = MINT_ACCENT
    gov_card.line.width = Pt(1)
    
    tf_gc = gov_card.text_frame
    tf_gc.word_wrap = True
    p_gh = tf_gc.paragraphs[0]
    p_gh.text = "GOVERNANCE & SAFETY STANDARD:"
    p_gh.font.size = Pt(10)
    p_gh.font.bold = True
    p_gh.font.color.rgb = MINT_ACCENT
    p_gh.font.name = 'Segoe UI'
    
    p_gd = tf_gc.add_paragraph()
    p_gd.text = "\"Designed with healthcare safety, deterministic emergency detection, and controlled human clinical oversight at every operational stage.\""
    p_gd.font.size = Pt(9)
    p_gd.font.color.rgb = TEXT_WHITE
    p_gd.font.name = 'Segoe UI'
    p_gd.space_before = Pt(2)

    # Right: Real Audit Trail & Trust Center Screenshot
    sc8a = os.path.join(SCREENSHOTS_DIR, "09_institutional_audit_trail.png")
    if os.path.exists(sc8a):
        s8.shapes.add_picture(sc8a, Inches(6.8), Inches(1.45), Inches(5.733), Inches(2.7))

    sc8b = os.path.join(SCREENSHOTS_DIR, "10_trust_center_governance.png")
    if os.path.exists(sc8b):
        s8.shapes.add_picture(sc8b, Inches(6.8), Inches(4.3), Inches(5.733), Inches(2.7))

    # =========================================================================
    # SLIDE 9 — IMPACT, VALIDATION & FUTURE SCOPE
    # =========================================================================
    s9 = prs.slides.add_slide(blank_layout)
    set_slide_background(s9)
    add_header(s9, "Impact, Empirical Validation & Future Scope", "IMPACT & ROADMAP", 9)

    # 3 Columns
    # Section 1: Expected Impact
    s1_card = s9.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.45), Inches(3.7), Inches(5.5))
    s1_card.fill.solid()
    s1_card.fill.fore_color.rgb = CARD_BG
    s1_card.line.color.rgb = BLUE_GLOW
    s1_card.line.width = Pt(1.5)
    
    tf1 = s1_card.text_frame
    tf1.word_wrap = True
    p1_h = tf1.paragraphs[0]
    p1_h.text = "SECTION 1 — EXPECTED IMPACT"
    p1_h.font.size = Pt(11)
    p1_h.font.bold = True
    p1_h.font.color.rgb = BLUE_GLOW
    p1_h.font.name = 'Segoe UI'
    p1_h.space_after = Pt(10)
    
    impacts = [
        "Faster Patient Structuring: Rapidly organizes unstructured complaints into clinical timelines.",
        "Early Red-Flag Detection: Prioritizes potential cardiorespiratory and high-risk cases.",
        "Multimodal Flexibility: Accommodates text, voice, and medical lab reports.",
        "Seamless Handoff: Smooth data flow across Patient, Nurse, and Doctor.",
        "Complete Traceability: Full audit history with clinician sign-off."
    ]
    for imp in impacts:
        p = tf1.add_paragraph()
        p.text = f"✔  {imp}"
        p.font.size = Pt(9)
        p.font.color.rgb = TEXT_LIGHT
        p.font.name = 'Segoe UI'
        p.space_before = Pt(4)

    # Section 2: Real Project Validation
    s2_card = s9.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(4.8), Inches(1.45), Inches(3.7), Inches(5.5))
    s2_card.fill.solid()
    s2_card.fill.fore_color.rgb = CARD_BG
    s2_card.line.color.rgb = MINT_ACCENT
    s2_card.line.width = Pt(1.5)
    
    tf2 = s2_card.text_frame
    tf2.word_wrap = True
    p2_h = tf2.paragraphs[0]
    p2_h.text = "SECTION 2 — REAL VALIDATION"
    p2_h.font.size = Pt(11)
    p2_h.font.bold = True
    p2_h.font.color.rgb = MINT_ACCENT
    p2_h.font.name = 'Segoe UI'
    p2_h.space_after = Pt(8)
    
    p2_sub = tf2.add_paragraph()
    p2_sub.text = "Empirical Backend Automation Test Results:"
    p2_sub.font.size = Pt(9.5)
    p2_sub.font.bold = True
    p2_sub.font.color.rgb = TEXT_WHITE
    p2_sub.font.name = 'Segoe UI'
    p2_sub.space_after = Pt(6)
    
    stats = [
        ("32", "Health AI functional & safety scenarios"),
        ("192", "Backend-connected executions"),
        ("192 Passed", "0 Failures recorded in test runner"),
        ("13 / 13", "E2E verification phases passed"),
        ("3,448 B", "Binary ReportLab referral PDF generated")
    ]
    for num_s, desc_s in stats:
        p_n = tf2.add_paragraph()
        p_n.text = f"• {num_s}: "
        p_n.font.size = Pt(10)
        p_n.font.bold = True
        p_n.font.color.rgb = TEAL_ACCENT
        p_n.font.name = 'Segoe UI'
        
        r_d = p_n.add_run()
        r_d.text = desc_s
        r_d.font.size = Pt(9)
        r_d.font.bold = False
        r_d.font.color.rgb = TEXT_LIGHT
        r_d.font.name = 'Segoe UI'
        p_n.space_before = Pt(3)

    # Section 3: Future Scope
    s3_card = s9.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.8), Inches(1.45), Inches(3.733), Inches(5.5))
    s3_card.fill.solid()
    s3_card.fill.fore_color.rgb = CARD_BG
    s3_card.line.color.rgb = RGBColor(168, 85, 247)
    s3_card.line.width = Pt(1.5)
    
    tf3 = s3_card.text_frame
    tf3.word_wrap = True
    p3_h = tf3.paragraphs[0]
    p3_h.text = "SECTION 3 — FUTURE SCOPE"
    p3_h.font.size = Pt(11)
    p3_h.font.bold = True
    p3_h.font.color.rgb = RGBColor(168, 85, 247)
    p3_h.font.name = 'Segoe UI'
    p3_h.space_after = Pt(10)
    
    futures = [
        "Wider Multilingual Coverage (Expanded regional Indian dialect models).",
        "Hospital EHR / ABDM Integration (FHIR-compliant patient records).",
        "Expanded Clinical Validation (Multi-center hospital pilot trials).",
        "Offline Low-Connectivity Sync (Edge SQLite replication for rural PHCs).",
        "Wearable Vitals Streaming (Continuous IoT telemetry ingestion)."
    ]
    for fut in futures:
        p = tf3.add_paragraph()
        p.text = f"➔  {fut}"
        p.font.size = Pt(9)
        p.font.color.rgb = TEXT_LIGHT
        p.font.name = 'Segoe UI'
        p.space_before = Pt(4)

    # =========================================================================
    # SLIDE 10 — THANK YOU
    # =========================================================================
    s10 = prs.slides.add_slide(blank_layout)
    set_slide_background(s10)

    # Decorative header bar
    bar10 = s10.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.9), Inches(11.733), Inches(0.06))
    bar10.fill.solid()
    bar10.fill.fore_color.rgb = TEAL_ACCENT
    bar10.line.fill.background()

    # Thank You Main Text
    t10 = s10.shapes.add_textbox(Inches(0.8), Inches(1.4), Inches(11.733), Inches(1.8))
    tf10 = t10.text_frame
    tf10.word_wrap = True
    
    p_ty = tf10.paragraphs[0]
    p_ty.text = "THANK YOU"
    p_ty.font.size = Pt(48)
    p_ty.font.bold = True
    p_ty.font.color.rgb = TEXT_WHITE
    p_ty.font.name = 'Segoe UI'
    
    p_q = tf10.add_paragraph()
    p_q.text = "Questions & Demonstration"
    p_q.font.size = Pt(20)
    p_q.font.color.rgb = TEAL_ACCENT
    p_q.font.name = 'Segoe UI'
    p_q.space_before = Pt(4)

    # Card with Editable Fields
    fcard10 = s10.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(3.4), Inches(6.5), Inches(3.4))
    fcard10.fill.solid()
    fcard10.fill.fore_color.rgb = CARD_BG
    fcard10.line.color.rgb = CARD_BORDER
    fcard10.line.width = Pt(1)
    
    tf_f10 = fcard10.text_frame
    tf_f10.word_wrap = True
    
    p_proj = tf_f10.paragraphs[0]
    p_proj.text = "PROJECT:  SWASTHYA TRIAGE"
    p_proj.font.size = Pt(14)
    p_proj.font.bold = True
    p_proj.font.color.rgb = BLUE_GLOW
    p_proj.font.name = 'Segoe UI'
    p_proj.space_after = Pt(12)
    
    ty_fields = [
        ("Team Name:", "[ADD]"),
        ("Problem Statement ID:", "[ADD]"),
        ("Live Website / QR Code:", "[ADD LATER]"),
        ("Technology Stack:", "React 19 + TypeScript + FastAPI + SQLite")
    ]
    for idx, (lbl, val) in enumerate(ty_fields):
        p = tf_f10.add_paragraph()
        p.text = f"{lbl} "
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = TEAL_ACCENT if idx < 3 else MINT_ACCENT
        p.font.name = 'Segoe UI'
        
        r = p.add_run()
        r.text = val
        r.font.size = Pt(11)
        r.font.bold = False
        r.font.color.rgb = RGBColor(253, 224, 71) if idx < 3 else TEXT_WHITE
        r.font.name = 'Segoe UI'
        p.space_after = Pt(8)

    # Right Card: Closing Highlight Summary
    rcard10 = s10.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.6), Inches(3.4), Inches(4.933), Inches(3.4))
    rcard10.fill.solid()
    rcard10.fill.fore_color.rgb = ACCENT_CARD
    rcard10.line.color.rgb = BLUE_GLOW
    rcard10.line.width = Pt(1.5)
    
    tf_rc = rcard10.text_frame
    tf_rc.word_wrap = True
    p_rch = tf_rc.paragraphs[0]
    p_rch.text = "Swasthya Triage Summary:"
    p_rch.font.size = Pt(13)
    p_rch.font.bold = True
    p_rch.font.color.rgb = TEAL_ACCENT
    p_rch.font.name = 'Segoe UI'
    p_rch.space_after = Pt(8)
    
    closing_points = [
        "100% Real Full-Stack Implementation",
        "Multimodal Intake (Text, Voice, OCR)",
        "Strictly Non-Diagnostic Clinician-Supervised AI",
        "Complete RBAC for Patient, Nurse, and Doctor",
        "192/192 Backend Automated Safety Tests Passed"
    ]
    for pt in closing_points:
        p = tf_rc.add_paragraph()
        p.text = f"★  {pt}"
        p.font.size = Pt(10)
        p.font.color.rgb = TEXT_WHITE
        p.font.name = 'Segoe UI'
        p.space_before = Pt(4)

    # Save presentation
    prs.save(OUTPUT_PPTX)
    print(f"Presentation generated successfully: {OUTPUT_PPTX}")

if __name__ == "__main__":
    create_presentation()
