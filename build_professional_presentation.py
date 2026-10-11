import os
import sys
import pptx
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

SCREENSHOTS_DIR = os.path.join(os.path.dirname(__file__), "presentation_screenshots")
OUTPUT_PPTX = os.path.join(os.path.dirname(__file__), "Swasthya_Triage_Hackathon_Presentation_Professional.pptx")

# ==============================================================================
# PROFESSIONAL HEALTHCARE AI PALETTE
# ==============================================================================
COLOR_NAVY_DARK    = RGBColor(11, 18, 32)      # #0B1220 - Dark Hero Background
COLOR_NAVY_SURFACE = RGBColor(20, 36, 58)      # #14243A - Dark Card Surface
COLOR_NAVY_BORDER  = RGBColor(34, 58, 88)      # #223A58 - Dark Card Border
COLOR_LIGHT_BG     = RGBColor(244, 248, 251)   # #F4F8FB - Clean Clinical Background
COLOR_WHITE        = RGBColor(255, 255, 255)   # #FFFFFF - Pure White Surface
COLOR_BORDER_LIGHT = RGBColor(218, 228, 237)   # #DAE4ED - Light Card Border
COLOR_BORDER_TEAL  = RGBColor(19, 184, 166)    # #13B8A6 - Primary Teal Accent
COLOR_TEAL_DARK    = RGBColor(15, 118, 110)    # #0F766E - Deep Teal Text
COLOR_CYAN_LIGHT   = RGBColor(115, 224, 208)   # #73E0D0 - Soft Cyan
COLOR_MINT         = RGBColor(16, 185, 129)    # #10B981 - Success Mint
COLOR_RED_ALERT    = RGBColor(217, 72, 72)     # #D94848 - Emergency Red Accent
COLOR_AMBER        = RGBColor(217, 119, 6)     # #D97706 - Warning / Placeholder
COLOR_BLUE_SKY     = RGBColor(14, 165, 233)    # #0EA5E9 - Sky Blue
COLOR_PURPLE       = RGBColor(168, 85, 247)    # #A855F7 - Purple Accent

# Typography Colors
COLOR_TEXT_PRIMARY_DARK = RGBColor(15, 23, 42)    # #0F172A - Deep Slate Header
COLOR_TEXT_SECONDARY    = RGBColor(82, 97, 116)   # #526174 - Secondary Text
COLOR_TEXT_MUTED        = RGBColor(148, 163, 184) # #94A3B8 - Caption Text
COLOR_TEXT_WHITE        = RGBColor(255, 255, 255) # #FFFFFF - Hero White

FONT_HEADING = "Segoe UI"
FONT_BODY    = "Segoe UI"

def create_deck():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    def add_slide_bg(slide, is_dark=False):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
        bg.fill.solid()
        bg.fill.fore_color.rgb = COLOR_NAVY_DARK if is_dark else COLOR_LIGHT_BG
        bg.line.fill.background()
        return bg

    def add_header(slide, title, category, slide_num, is_dark=False):
        # Category Pill
        pill = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.38), Inches(4.8), Inches(0.32))
        pill.fill.solid()
        pill.fill.fore_color.rgb = RGBColor(204, 251, 241) if not is_dark else RGBColor(19, 78, 74)
        pill.line.color.rgb = COLOR_BORDER_TEAL
        pill.line.width = Pt(1)
        tf_p = pill.text_frame
        tf_p.word_wrap = True
        tf_p.margin_left = Inches(0.15)
        tf_p.margin_right = Inches(0.15)
        tf_p.margin_top = Inches(0.04)
        tf_p.margin_bottom = Inches(0.04)
        p_p = tf_p.paragraphs[0]
        p_p.text = f"  {category.upper()}"
        p_p.font.size = Pt(9.5)
        p_p.font.bold = True
        p_p.font.color.rgb = COLOR_TEAL_DARK if not is_dark else COLOR_CYAN_LIGHT
        p_p.font.name = FONT_HEADING
        p_p.alignment = PP_ALIGN.LEFT

        # Slide Number
        num_box = slide.shapes.add_textbox(Inches(11.333), Inches(0.38), Inches(1.2), Inches(0.35))
        tf_n = num_box.text_frame
        tf_n.margin_left = Inches(0)
        tf_n.margin_top = Inches(0)
        p_n = tf_n.paragraphs[0]
        p_n.text = f"{slide_num:02d} / 10"
        p_n.font.size = Pt(11)
        p_n.font.bold = True
        p_n.font.color.rgb = COLOR_TEXT_MUTED if not is_dark else COLOR_CYAN_LIGHT
        p_n.font.name = FONT_HEADING
        p_n.alignment = PP_ALIGN.RIGHT

        # Main Slide Title
        t_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.72), Inches(11.733), Inches(0.55))
        tf_t = t_box.text_frame
        tf_t.word_wrap = True
        tf_t.margin_left = Inches(0)
        tf_t.margin_top = Inches(0)
        p_t = tf_t.paragraphs[0]
        p_t.text = title
        p_t.font.size = Pt(20)
        p_t.font.bold = True
        p_t.font.color.rgb = COLOR_TEXT_WHITE if is_dark else COLOR_TEXT_PRIMARY_DARK
        p_t.font.name = FONT_HEADING
        p_t.alignment = PP_ALIGN.LEFT

    # =========================================================================
    # SLIDE 1 — COVER (HERO DARK)
    # =========================================================================
    s1 = prs.slides.add_slide(blank_layout)
    add_slide_bg(s1, is_dark=True)

    top_bar = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(0.7), Inches(11.733), Inches(0.05))
    top_bar.fill.solid()
    top_bar.fill.fore_color.rgb = COLOR_BORDER_TEAL
    top_bar.line.fill.background()

    pill1 = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.95), Inches(5.8), Inches(0.34))
    pill1.fill.solid()
    pill1.fill.fore_color.rgb = RGBColor(19, 78, 74)
    pill1.line.color.rgb = COLOR_BORDER_TEAL
    pill1.line.width = Pt(1)
    tf_p1 = pill1.text_frame
    tf_p1.margin_left = Inches(0.15)
    tf_p1.margin_top = Inches(0.04)
    p_p1 = tf_p1.paragraphs[0]
    p_p1.text = "  CLINICAL DECISION SUPPORT & MULTIMODAL TRIAGE"
    p_p1.font.size = Pt(9.5)
    p_p1.font.bold = True
    p_p1.font.color.rgb = COLOR_CYAN_LIGHT
    p_p1.font.name = FONT_HEADING
    p_p1.alignment = PP_ALIGN.LEFT

    t1_box = s1.shapes.add_textbox(Inches(0.8), Inches(1.4), Inches(11.733), Inches(1.3))
    tf1 = t1_box.text_frame
    tf1.word_wrap = True
    tf1.margin_left = Inches(0)
    tf1.margin_top = Inches(0)
    
    p1 = tf1.paragraphs[0]
    p1.text = "SWASTHYA TRIAGE"
    p1.font.size = Pt(36)
    p1.font.bold = True
    p1.font.color.rgb = COLOR_TEXT_WHITE
    p1.font.name = FONT_HEADING
    p1.alignment = PP_ALIGN.LEFT

    p1_sub = tf1.add_paragraph()
    p1_sub.text = "Clinician-Supervised Multimodal Patient Intake & Safe Triage Assistance"
    p1_sub.font.size = Pt(14)
    p1_sub.font.color.rgb = COLOR_CYAN_LIGHT
    p1_sub.font.name = FONT_HEADING
    p1_sub.alignment = PP_ALIGN.LEFT
    p1_sub.space_before = Pt(4)

    card1 = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(2.85), Inches(5.8), Inches(3.8))
    card1.fill.solid()
    card1.fill.fore_color.rgb = COLOR_NAVY_SURFACE
    card1.line.color.rgb = COLOR_NAVY_BORDER
    card1.line.width = Pt(1)

    tf_c1 = card1.text_frame
    tf_c1.word_wrap = True
    tf_c1.vertical_anchor = MSO_ANCHOR.TOP
    tf_c1.margin_left = Inches(0.3)
    tf_c1.margin_top = Inches(0.3)
    tf_c1.margin_right = Inches(0.3)
    tf_c1.margin_bottom = Inches(0.3)
    
    placeholders = [
        ("Problem Statement:", "[ADD LATER]"),
        ("Problem Statement ID:", "[ADD LATER]"),
        ("Team Name:", "[ADD LATER]"),
        ("Domain:", "[ADD LATER]"),
        ("Tagline:", "[ADD LATER]"),
        ("System Architecture:", "FastAPI + Python + React 19 + AI Safety Pipeline")
    ]
    for idx, (lbl, val) in enumerate(placeholders):
        p_row = tf_c1.paragraphs[0] if idx == 0 else tf_c1.add_paragraph()
        p_row.text = f"{lbl} "
        p_row.font.size = Pt(11)
        p_row.font.bold = True
        p_row.font.color.rgb = COLOR_CYAN_LIGHT if idx < 5 else COLOR_MINT
        p_row.font.name = FONT_HEADING
        p_row.alignment = PP_ALIGN.LEFT
        
        run_v = p_row.add_run()
        run_v.text = val
        run_v.font.size = Pt(11)
        run_v.font.bold = False
        run_v.font.color.rgb = RGBColor(253, 224, 71) if idx < 5 else COLOR_TEXT_WHITE
        run_v.font.name = FONT_HEADING
        p_row.space_after = Pt(8)

    sc1 = os.path.join(SCREENSHOTS_DIR, "01_patient_login.png")
    if os.path.exists(sc1):
        frame1 = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(2.85), Inches(5.733), Inches(3.8))
        frame1.fill.solid()
        frame1.fill.fore_color.rgb = COLOR_NAVY_SURFACE
        frame1.line.color.rgb = COLOR_NAVY_BORDER
        frame1.line.width = Pt(1)
        
        cap1 = s1.shapes.add_textbox(Inches(6.9), Inches(2.9), Inches(5.5), Inches(0.35))
        tf_cap = cap1.text_frame
        tf_cap.margin_left = Inches(0)
        tf_cap.margin_top = Inches(0)
        p_cap = tf_cap.paragraphs[0]
        p_cap.text = "LIVE PLATFORM • PATIENT & CLINICAL AUTHENTICATION"
        p_cap.font.size = Pt(9)
        p_cap.font.bold = True
        p_cap.font.color.rgb = COLOR_CYAN_LIGHT
        p_cap.font.name = FONT_HEADING
        p_cap.alignment = PP_ALIGN.LEFT
        
        s1.shapes.add_picture(sc1, Inches(6.9), Inches(3.25), Inches(5.533), Inches(3.3))

    bbar = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(6.85), Inches(11.733), Inches(0.35))
    bbar.fill.solid()
    bbar.fill.fore_color.rgb = RGBColor(15, 27, 48)
    bbar.line.fill.background()
    p_bb = bbar.text_frame.paragraphs[0]
    p_bb.text = "TECH STACK:  React 19  •  TypeScript  •  FastAPI  •  Python 3.12  •  SQLite / PostgreSQL  •  ReportLab PDF"
    p_bb.font.size = Pt(9)
    p_bb.font.bold = True
    p_bb.font.color.rgb = COLOR_TEXT_MUTED
    p_bb.font.name = FONT_HEADING
    p_bb.alignment = PP_ALIGN.CENTER

    # =========================================================================
    # SLIDE 2 — PROBLEM & MOTIVATION (CLEAN CLINICAL CANVAS)
    # =========================================================================
    s2 = prs.slides.add_slide(blank_layout)
    add_slide_bg(s2, is_dark=False)
    add_header(s2, "The Problem: Unstructured Triage & Information Overload", "PROBLEM & MOTIVATION", 2)

    cards_p = [
        ("1. Unstructured Patient Complaints", "Patients present symptoms in colloquial, disorganized narratives across text, speech, and paper reports where critical red-flags are easily obscured."),
        ("2. Delayed Urgency Prioritization", "Emergency and outpatient queues struggle to rapidly distinguish routine complaints from time-sensitive emergencies prior to formal bedside assessment."),
        ("3. Need for Clinician-Supervised AI", "Generic chatbots hallucinate and provide unsupervised advice. Clinical environments require non-diagnostic AI with mandatory physician oversight.")
    ]
    y_p = 1.35
    for title_c, desc_c in cards_p:
        p_card = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(y_p), Inches(5.7), Inches(1.3))
        p_card.fill.solid()
        p_card.fill.fore_color.rgb = COLOR_WHITE
        p_card.line.color.rgb = COLOR_BORDER_LIGHT
        p_card.line.width = Pt(1)

        tf = p_card.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.TOP
        tf.margin_left = Inches(0.25)
        tf.margin_top = Inches(0.18)
        tf.margin_right = Inches(0.25)
        tf.margin_bottom = Inches(0.18)

        p_t = tf.paragraphs[0]
        p_t.text = title_c
        p_t.font.size = Pt(11.5)
        p_t.font.bold = True
        p_t.font.color.rgb = COLOR_TEXT_PRIMARY_DARK
        p_t.font.name = FONT_HEADING
        p_t.alignment = PP_ALIGN.LEFT

        p_d = tf.add_paragraph()
        p_d.text = desc_c
        p_d.font.size = Pt(9.5)
        p_d.font.color.rgb = COLOR_TEXT_SECONDARY
        p_d.font.name = FONT_BODY
        p_d.alignment = PP_ALIGN.LEFT
        p_d.space_before = Pt(3)
        y_p += 1.4

    flow_card = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(5.6), Inches(5.7), Inches(1.5))
    flow_card.fill.solid()
    flow_card.fill.fore_color.rgb = RGBColor(241, 245, 249)
    flow_card.line.color.rgb = COLOR_BORDER_TEAL
    flow_card.line.width = Pt(1)

    tf_fc = flow_card.text_frame
    tf_fc.word_wrap = True
    tf_fc.vertical_anchor = MSO_ANCHOR.TOP
    tf_fc.margin_left = Inches(0.25)
    tf_fc.margin_top = Inches(0.18)
    tf_fc.margin_right = Inches(0.25)
    tf_fc.margin_bottom = Inches(0.18)

    p_trad = tf_fc.paragraphs[0]
    p_trad.text = "TRADITIONAL WORKFLOW:"
    p_trad.font.size = Pt(9.5)
    p_trad.font.bold = True
    p_trad.font.color.rgb = COLOR_RED_ALERT
    p_trad.font.name = FONT_HEADING
    p_trad.alignment = PP_ALIGN.LEFT

    p_tf = tf_fc.add_paragraph()
    p_tf.text = "Patient  →  Unstructured Info  →  Manual Interpretation  →  Delayed Prioritization"
    p_tf.font.size = Pt(9)
    p_tf.font.color.rgb = COLOR_TEXT_SECONDARY
    p_tf.font.name = FONT_BODY
    p_tf.alignment = PP_ALIGN.LEFT
    p_tf.space_after = Pt(6)

    p_app = tf_fc.add_paragraph()
    p_app.text = "SWASTHYA TRIAGE APPROACH:"
    p_app.font.size = Pt(9.5)
    p_app.font.bold = True
    p_app.font.color.rgb = COLOR_TEAL_DARK
    p_app.font.name = FONT_HEADING
    p_app.alignment = PP_ALIGN.LEFT

    p_af = tf_fc.add_paragraph()
    p_af.text = "Patient  →  Structured Multimodal Intake  →  AI Safety & Triage  →  Human Clinical Review"
    p_af.font.size = Pt(9)
    p_af.font.bold = True
    p_af.font.color.rgb = COLOR_TEXT_PRIMARY_DARK
    p_af.font.name = FONT_BODY
    p_af.alignment = PP_ALIGN.LEFT

    sc2 = os.path.join(SCREENSHOTS_DIR, "02_patient_portal_dashboard.png")
    if os.path.exists(sc2):
        s2.shapes.add_picture(sc2, Inches(6.8), Inches(1.35), Inches(5.733), Inches(5.75))

    # =========================================================================
    # SLIDE 3 — OUR SOLUTION (CONNECTED 6-STAGE WORKFLOW)
    # =========================================================================
    s3 = prs.slides.add_slide(blank_layout)
    add_slide_bg(s3, is_dark=False)
    add_header(s3, "Introducing Swasthya Triage: 6-Stage Clinical Handoff", "OUR SOLUTION", 3)

    banner3 = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.35), Inches(11.733), Inches(0.48))
    banner3.fill.solid()
    banner3.fill.fore_color.rgb = RGBColor(204, 251, 241)
    banner3.line.color.rgb = COLOR_BORDER_TEAL
    banner3.line.width = Pt(1)
    tf_b3 = banner3.text_frame
    tf_b3.margin_top = Inches(0.08)
    p_b3 = tf_b3.paragraphs[0]
    p_b3.text = "  CORE PRINCIPLE:  \"AI assists the workflow. Clinicians make the final decision.\""
    p_b3.font.size = Pt(11)
    p_b3.font.bold = True
    p_b3.font.color.rgb = COLOR_TEAL_DARK
    p_b3.font.name = FONT_HEADING
    p_b3.alignment = PP_ALIGN.LEFT

    stages = [
        ("01. Patient Intake", "Multimodal input via Text, Voice, or Lab OCR"),
        ("02. Health AI", "Symptom extraction & red-flag evaluation"),
        ("03. Priority Triage", "Urgency signal calculation with evidence"),
        ("04. Nurse Queue", "Prioritized triaging & bedside vitals"),
        ("05. Reviewer MD", "Clinical confirmation, notes, or override"),
        ("06. Referral & Audit", "Instant PDF generation & audit trail")
    ]
    w_box = 1.83
    x_box = 0.8
    for idx, (st_t, st_d) in enumerate(stages):
        box_s = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x_box), Inches(1.95), Inches(w_box), Inches(1.4))
        box_s.fill.solid()
        box_s.fill.fore_color.rgb = COLOR_WHITE
        box_s.line.color.rgb = COLOR_BORDER_TEAL if idx in [1, 2] else COLOR_BORDER_LIGHT
        box_s.line.width = Pt(1.2)

        tf_s = box_s.text_frame
        tf_s.word_wrap = True
        tf_s.vertical_anchor = MSO_ANCHOR.TOP
        tf_s.margin_left = Inches(0.15)
        tf_s.margin_top = Inches(0.15)
        tf_s.margin_right = Inches(0.15)
        tf_s.margin_bottom = Inches(0.15)

        p_st = tf_s.paragraphs[0]
        p_st.text = st_t
        p_st.font.size = Pt(10)
        p_st.font.bold = True
        p_st.font.color.rgb = COLOR_TEAL_DARK if idx in [1, 2] else COLOR_TEXT_PRIMARY_DARK
        p_st.font.name = FONT_HEADING
        p_st.alignment = PP_ALIGN.LEFT

        p_sd = tf_s.add_paragraph()
        p_sd.text = st_d
        p_sd.font.size = Pt(8.5)
        p_sd.font.color.rgb = COLOR_TEXT_SECONDARY
        p_sd.font.name = FONT_BODY
        p_sd.alignment = PP_ALIGN.LEFT
        p_sd.space_before = Pt(3)

        x_box += 1.98

    cap_card = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(3.5), Inches(4.3), Inches(3.6))
    cap_card.fill.solid()
    cap_card.fill.fore_color.rgb = COLOR_WHITE
    cap_card.line.color.rgb = COLOR_BORDER_LIGHT
    cap_card.line.width = Pt(1)

    tf_cc = cap_card.text_frame
    tf_cc.word_wrap = True
    tf_cc.vertical_anchor = MSO_ANCHOR.TOP
    tf_cc.margin_left = Inches(0.25)
    tf_cc.margin_top = Inches(0.2)
    tf_cc.margin_right = Inches(0.25)
    tf_cc.margin_bottom = Inches(0.2)

    p_cch = tf_cc.paragraphs[0]
    p_cch.text = "Implemented Core Capabilities:"
    p_cch.font.size = Pt(11.5)
    p_cch.font.bold = True
    p_cch.font.color.rgb = COLOR_TEXT_PRIMARY_DARK
    p_cch.font.name = FONT_HEADING
    p_cch.alignment = PP_ALIGN.LEFT

    caps = [
        "Patient Self-Intake & JWT Authentication",
        "Conversational Health AI with Safety Guardrails",
        "Deterministic Urgency Rules (Cardiorespiratory, etc.)",
        "Role-Based Access Control (Patient, Nurse, Doctor, Admin)",
        "Nurse Bedside Vitals (Temp, BP, SpO2, Heart Rate)",
        "Doctor Review & Direct Override Protocol",
        "Referral Generation with ReportLab PDF Engine"
    ]
    for c in caps:
        p_c = tf_cc.add_paragraph()
        p_c.text = f"✔  {c}"
        p_c.font.size = Pt(9)
        p_c.font.color.rgb = COLOR_TEXT_SECONDARY
        p_c.font.name = FONT_BODY
        p_c.alignment = PP_ALIGN.LEFT
        p_c.space_before = Pt(3)

    sc3a = os.path.join(SCREENSHOTS_DIR, "04_patient_intake_multimodal.png")
    if os.path.exists(sc3a):
        s3.shapes.add_picture(sc3a, Inches(5.25), Inches(3.5), Inches(3.85), Inches(3.6))

    sc3b = os.path.join(SCREENSHOTS_DIR, "08_medical_reviewer_workspace.png")
    if os.path.exists(sc3b):
        s3.shapes.add_picture(sc3b, Inches(9.25), Inches(3.5), Inches(3.283), Inches(3.6))

    # =========================================================================
    # SLIDE 4 — AI & SAFETY ENGINE (DIAGRAM-LED PIPELINE)
    # =========================================================================
    s4 = prs.slides.add_slide(blank_layout)
    add_slide_bg(s4, is_dark=False)
    add_header(s4, "Health AI — From Patient Input to Safe Assistance", "AI & SAFETY ENGINE", 4)

    ai_pipeline = [
        ("1. Multimodal Patient Input", "Accepts natural language text, voice recording, or uploaded lab reports."),
        ("2. Clinical Structuring", "Extracts Chief Complaint, Symptoms, Timeline, and Medical History."),
        ("3. Safety & Red-Flag Rules", "Evaluates deterministic emergency rules (e.g. Cardiorespiratory distress)."),
        ("4. Triage Prioritization", "Computes Urgency Level (Emergency, Urgent, Routine) with evidence citations."),
        ("5. Guardrailed Non-Diagnostic Output", "Delivers structured summary with mandatory clinician review protocol.")
    ]
    y_step = 1.35
    for s_title, s_desc in ai_pipeline:
        s_box = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(y_step), Inches(5.6), Inches(1.0))
        s_box.fill.solid()
        s_box.fill.fore_color.rgb = COLOR_WHITE
        s_box.line.color.rgb = COLOR_RED_ALERT if "Safety" in s_title else (COLOR_BORDER_TEAL if "Guardrailed" in s_title else COLOR_BORDER_LIGHT)
        s_box.line.width = Pt(1.2)

        tf_s = s_box.text_frame
        tf_s.word_wrap = True
        tf_s.vertical_anchor = MSO_ANCHOR.TOP
        tf_s.margin_left = Inches(0.2)
        tf_s.margin_top = Inches(0.12)
        tf_s.margin_right = Inches(0.2)
        tf_s.margin_bottom = Inches(0.12)

        p_st = tf_s.paragraphs[0]
        p_st.text = s_title
        p_st.font.size = Pt(10.5)
        p_st.font.bold = True
        p_st.font.color.rgb = COLOR_RED_ALERT if "Safety" in s_title else (COLOR_TEAL_DARK if "Guardrailed" in s_title else COLOR_TEXT_PRIMARY_DARK)
        p_st.font.name = FONT_HEADING
        p_st.alignment = PP_ALIGN.LEFT

        p_sd = tf_s.add_paragraph()
        p_sd.text = s_desc
        p_sd.font.size = Pt(9)
        p_sd.font.color.rgb = COLOR_TEXT_SECONDARY
        p_sd.font.name = FONT_BODY
        p_sd.alignment = PP_ALIGN.LEFT
        y_step += 1.08

    sc4 = os.path.join(SCREENSHOTS_DIR, "03_health_ai_chat.png")
    if os.path.exists(sc4):
        s4.shapes.add_picture(sc4, Inches(6.7), Inches(1.35), Inches(5.833), Inches(4.3))

    safe_box = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.7), Inches(5.8), Inches(5.833), Inches(1.3))
    safe_box.fill.solid()
    safe_box.fill.fore_color.rgb = RGBColor(254, 242, 242)
    safe_box.line.color.rgb = COLOR_RED_ALERT
    safe_box.line.width = Pt(1.2)

    tf_sb = safe_box.text_frame
    tf_sb.word_wrap = True
    tf_sb.vertical_anchor = MSO_ANCHOR.TOP
    tf_sb.margin_left = Inches(0.25)
    tf_sb.margin_top = Inches(0.15)
    tf_sb.margin_right = Inches(0.25)
    tf_sb.margin_bottom = Inches(0.15)

    p_sbh = tf_sb.paragraphs[0]
    p_sbh.text = "🛡️ NON-DIAGNOSTIC SAFETY GUARANTEE:"
    p_sbh.font.size = Pt(10)
    p_sbh.font.bold = True
    p_sbh.font.color.rgb = COLOR_RED_ALERT
    p_sbh.font.name = FONT_HEADING
    p_sbh.alignment = PP_ALIGN.LEFT

    p_sbd = tf_sb.add_paragraph()
    p_sbd.text = "\"The AI identifies urgency signals and structures information. It does not independently diagnose conditions or prescribe medications.\""
    p_sbd.font.size = Pt(9.5)
    p_sbd.font.color.rgb = COLOR_TEXT_PRIMARY_DARK
    p_sbd.font.name = FONT_BODY
    p_sbd.alignment = PP_ALIGN.LEFT
    p_sbd.space_before = Pt(2)

    # =========================================================================
    # SLIDE 5 — MULTIMODAL & MULTILINGUAL HEALTH AI
    # =========================================================================
    s5 = prs.slides.add_slide(blank_layout)
    add_slide_bg(s5, is_dark=False)
    add_header(s5, "Multimodal & Multilingual Health AI: Unified Context", "MULTIMODAL INTAKE", 5)

    modes = [
        ("1. TEXT INTAKE", "Natural language text in English/Hindi with timeline & symptom capture.", COLOR_BORDER_TEAL),
        ("2. VOICE ASSIST", "Speech-to-text audio recording for elderly and accessibility-first patients.", COLOR_BLUE_SKY),
        ("3. MEDICAL OCR", "Lab report & prescription upload with OCR extraction of vitals/findings.", COLOR_MINT),
        ("4. CLINICAL IMAGES", "Visual symptom attachment analyzed alongside patient clinical history.", COLOR_PURPLE)
    ]
    w_m = 2.76
    x_m = 0.8
    for m_t, m_d, m_col in modes:
        m_box = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x_m), Inches(1.35), Inches(w_m), Inches(1.4))
        m_box.fill.solid()
        m_box.fill.fore_color.rgb = COLOR_WHITE
        m_box.line.color.rgb = m_col
        m_box.line.width = Pt(1.2)

        tf_m = m_box.text_frame
        tf_m.word_wrap = True
        tf_m.vertical_anchor = MSO_ANCHOR.TOP
        tf_m.margin_left = Inches(0.18)
        tf_m.margin_top = Inches(0.15)
        tf_m.margin_right = Inches(0.18)
        tf_m.margin_bottom = Inches(0.15)

        p_mt = tf_m.paragraphs[0]
        p_mt.text = m_t
        p_mt.font.size = Pt(10.5)
        p_mt.font.bold = True
        p_mt.font.color.rgb = m_col
        p_mt.font.name = FONT_HEADING
        p_mt.alignment = PP_ALIGN.LEFT

        p_md = tf_m.add_paragraph()
        p_md.text = m_d
        p_md.font.size = Pt(8.5)
        p_md.font.color.rgb = COLOR_TEXT_SECONDARY
        p_md.font.name = FONT_BODY
        p_md.alignment = PP_ALIGN.LEFT
        p_md.space_before = Pt(3)

        x_m += 2.98

    pbar5 = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(2.9), Inches(11.733), Inches(0.45))
    pbar5.fill.solid()
    pbar5.fill.fore_color.rgb = RGBColor(241, 245, 249)
    pbar5.line.color.rgb = COLOR_BORDER_TEAL
    pbar5.line.width = Pt(1)
    p_pb5 = pbar5.text_frame.paragraphs[0]
    p_pb5.text = "  Text + Voice + Reports + Images  ──▶  Input Processing  ──▶  Unified Clinical Context  ──▶  AI Safety Pipeline  ──▶  Patient Response  "
    p_pb5.font.size = Pt(9.5)
    p_pb5.font.bold = True
    p_pb5.font.color.rgb = COLOR_TEAL_DARK
    p_pb5.font.name = FONT_HEADING
    p_pb5.alignment = PP_ALIGN.CENTER

    lang_card = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(3.5), Inches(5.7), Inches(3.6))
    lang_card.fill.solid()
    lang_card.fill.fore_color.rgb = COLOR_WHITE
    lang_card.line.color.rgb = COLOR_BORDER_LIGHT
    lang_card.line.width = Pt(1)

    tf_lc = lang_card.text_frame
    tf_lc.word_wrap = True
    tf_lc.vertical_anchor = MSO_ANCHOR.TOP
    tf_lc.margin_left = Inches(0.25)
    tf_lc.margin_top = Inches(0.2)
    tf_lc.margin_right = Inches(0.25)
    tf_lc.margin_bottom = Inches(0.2)

    p_lch = tf_lc.paragraphs[0]
    p_lch.text = "Supported Multilingual Coverage:"
    p_lch.font.size = Pt(11.5)
    p_lch.font.bold = True
    p_lch.font.color.rgb = COLOR_TEXT_PRIMARY_DARK
    p_lch.font.name = FONT_HEADING
    p_lch.alignment = PP_ALIGN.LEFT

    p_lcd = tf_lc.add_paragraph()
    p_lcd.text = "• Core Verified Languages: English, Hindi, Odia\n• Multilingual Pipeline Ready: Bengali, Telugu, Tamil, Kannada, Malayalam, Marathi, Gujarati, Punjabi\n• Features: Mixed-language intake & automatic language detection"
    p_lcd.font.size = Pt(9)
    p_lcd.font.color.rgb = COLOR_TEXT_SECONDARY
    p_lcd.font.name = FONT_BODY
    p_lcd.alignment = PP_ALIGN.LEFT
    p_lcd.space_before = Pt(4)
    p_lcd.space_after = Pt(10)

    p_callout = tf_lc.add_paragraph()
    p_callout.text = "💬  \"Language changes — clinical safety does not.\""
    p_callout.font.size = Pt(10.5)
    p_callout.font.bold = True
    p_callout.font.color.rgb = COLOR_TEAL_DARK
    p_callout.font.name = FONT_HEADING
    p_callout.alignment = PP_ALIGN.LEFT

    sc5 = os.path.join(SCREENSHOTS_DIR, "04_patient_intake_multimodal.png")
    if os.path.exists(sc5):
        s5.shapes.add_picture(sc5, Inches(6.8), Inches(3.5), Inches(5.733), Inches(3.6))

    # =========================================================================
    # SLIDE 6 — CLINICAL WORKFLOW & ROLE-BASED ACCESS
    # =========================================================================
    s6 = prs.slides.add_slide(blank_layout)
    add_slide_bg(s6, is_dark=False)
    add_header(s6, "From Patient Intake to Hospital Clinical Workflow", "CLINICAL WORKFLOW & RBAC", 6)

    flow_steps = [
        ("PATIENT INTAKE", "Submits multimodal symptoms & history via secure portal.", COLOR_BORDER_TEAL),
        ("AI-ASSISTED TRIAGE", "Extracts timeline, matches red-flag rules, computes urgency level.", COLOR_BLUE_SKY),
        ("NURSE QUEUE & BEDSIDE", "Nurse reviews prioritized list and records live vitals (BP, SpO2).", COLOR_MINT),
        ("MEDICAL REVIEWER (MD)", "Physician evaluates evidence, confirms triage, or executes override.", COLOR_AMBER),
        ("REFERRAL & CARE PATH", "Generates institutional referral PDF or schedules direct consult.", COLOR_PURPLE)
    ]
    y_fl = 1.35
    for fl_t, fl_d, fl_c in flow_steps:
        box_f = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(y_fl), Inches(5.3), Inches(0.88))
        box_f.fill.solid()
        box_f.fill.fore_color.rgb = COLOR_WHITE
        box_f.line.color.rgb = fl_c
        box_f.line.width = Pt(1.2)

        tf = box_f.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.TOP
        tf.margin_left = Inches(0.18)
        tf.margin_top = Inches(0.1)
        tf.margin_right = Inches(0.18)
        tf.margin_bottom = Inches(0.1)

        p_t = tf.paragraphs[0]
        p_t.text = fl_t
        p_t.font.size = Pt(10)
        p_t.font.bold = True
        p_t.font.color.rgb = fl_c
        p_t.font.name = FONT_HEADING
        p_t.alignment = PP_ALIGN.LEFT

        p_d = tf.add_paragraph()
        p_d.text = fl_d
        p_d.font.size = Pt(8.5)
        p_d.font.color.rgb = COLOR_TEXT_SECONDARY
        p_d.font.name = FONT_BODY
        p_d.alignment = PP_ALIGN.LEFT
        y_fl += 0.95

    resp_box = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(6.15), Inches(5.3), Inches(0.95))
    resp_box.fill.solid()
    resp_box.fill.fore_color.rgb = RGBColor(241, 245, 249)
    resp_box.line.color.rgb = COLOR_BORDER_LIGHT
    resp_box.line.width = Pt(1)

    tf_rb = resp_box.text_frame
    tf_rb.word_wrap = True
    tf_rb.vertical_anchor = MSO_ANCHOR.TOP
    tf_rb.margin_left = Inches(0.2)
    tf_rb.margin_top = Inches(0.12)
    tf_rb.margin_right = Inches(0.2)
    tf_rb.margin_bottom = Inches(0.12)

    p_ai_r = tf_rb.paragraphs[0]
    p_ai_r.text = "AI ROLE:  Assists • Structures • Flags • Summarizes"
    p_ai_r.font.size = Pt(9.5)
    p_ai_r.font.bold = True
    p_ai_r.font.color.rgb = COLOR_TEAL_DARK
    p_ai_r.font.name = FONT_HEADING
    p_ai_r.alignment = PP_ALIGN.LEFT

    p_cl_r = tf_rb.add_paragraph()
    p_cl_r.text = "CLINICIAN ROLE:  Reviews • Overrides • Prescribes • Decides"
    p_cl_r.font.size = Pt(9.5)
    p_cl_r.font.bold = True
    p_cl_r.font.color.rgb = COLOR_TEXT_PRIMARY_DARK
    p_cl_r.font.name = FONT_HEADING
    p_cl_r.alignment = PP_ALIGN.LEFT
    p_cl_r.space_before = Pt(3)

    sc6a = os.path.join(SCREENSHOTS_DIR, "06_nurse_triage_queue.png")
    if os.path.exists(sc6a):
        s6.shapes.add_picture(sc6a, Inches(6.35), Inches(1.35), Inches(6.183), Inches(2.75))

    sc6b = os.path.join(SCREENSHOTS_DIR, "07_nurse_bedside_assessment.png")
    if os.path.exists(sc6b):
        s6.shapes.add_picture(sc6b, Inches(6.35), Inches(4.3), Inches(6.183), Inches(2.8))

    # =========================================================================
    # SLIDE 7 — TECHNICAL ARCHITECTURE (FULL PROPORTIONAL CARDS)
    # =========================================================================
    s7 = prs.slides.add_slide(blank_layout)
    add_slide_bg(s7, is_dark=False)
    add_header(s7, "System Architecture: Robust Full-Stack Engineering", "TECHNICAL ARCHITECTURE", 7)

    arch_tiers = [
        ("1. FRONTEND LAYER", "React 19 + TypeScript + Vite", [
            ("Portals:", "Role-based routing for Patient, Nurse, Doctor, Admin"),
            ("UI Design:", "Glassmorphic clinical command tokens & scroll reveals"),
            ("Live Views:", "Real-time prioritized triage queue & bedside entry")
        ], COLOR_BORDER_TEAL),
        ("2. BACKEND API", "FastAPI + Python 3.12 (Async)", [
            ("Endpoints:", "RESTful APIs with strict Pydantic schema validation"),
            ("Security:", "JWT token authentication (HS256) & BCrypt hashing"),
            ("PDF Output:", "Binary ReportLab referral letter generation")
        ], COLOR_BLUE_SKY),
        ("3. AI & SAFETY ENGINE", "Provider Abstraction Layer", [
            ("LLM Core:", "Structured extraction of symptoms & timeline"),
            ("Safety Rules:", "Hardcoded deterministic emergency red-flag engine"),
            ("Fallback:", "Zero-cloud offline deterministic fallback pipeline")
        ], COLOR_MINT),
        ("4. DATABASE & AUDIT", "SQLite / PostgreSQL Ready", [
            ("ORM Layer:", "Async SQLAlchemy with full relational model"),
            ("Clinical Store:", "Persistent cases, consents, reviews, and vitals"),
            ("Audit Trail:", "Immutable audit log ensuring total provenance")
        ], COLOR_PURPLE)
    ]
    w_t = 2.76
    x_t = 0.8
    for t_title, t_tech, t_points, t_col in arch_tiers:
        t_card = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x_t), Inches(1.35), Inches(w_t), Inches(4.3))
        t_card.fill.solid()
        t_card.fill.fore_color.rgb = COLOR_WHITE
        t_card.line.color.rgb = t_col
        t_card.line.width = Pt(1.5)

        tf = t_card.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.TOP
        tf.margin_left = Inches(0.2)
        tf.margin_top = Inches(0.2)
        tf.margin_right = Inches(0.2)
        tf.margin_bottom = Inches(0.2)

        p_h = tf.paragraphs[0]
        p_h.text = t_title
        p_h.font.size = Pt(10.5)
        p_h.font.bold = True
        p_h.font.color.rgb = t_col
        p_h.font.name = FONT_HEADING
        p_h.alignment = PP_ALIGN.LEFT

        p_sub = tf.add_paragraph()
        p_sub.text = t_tech
        p_sub.font.size = Pt(9.5)
        p_sub.font.bold = True
        p_sub.font.color.rgb = COLOR_TEXT_PRIMARY_DARK
        p_sub.font.name = FONT_HEADING
        p_sub.alignment = PP_ALIGN.LEFT
        p_sub.space_after = Pt(8)

        for pt_lbl, pt_val in t_points:
            p_p = tf.add_paragraph()
            p_p.text = f"• {pt_lbl} "
            p_p.font.size = Pt(8.5)
            p_p.font.bold = True
            p_p.font.color.rgb = COLOR_TEXT_PRIMARY_DARK
            p_p.font.name = FONT_HEADING
            p_p.alignment = PP_ALIGN.LEFT
            
            r_pv = p_p.add_run()
            r_pv.text = pt_val
            r_pv.font.size = Pt(8.5)
            r_pv.font.bold = False
            r_pv.font.color.rgb = COLOR_TEXT_SECONDARY
            r_pv.font.name = FONT_BODY
            p_p.space_before = Pt(4)

        x_t += 2.98

    achain = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(5.8), Inches(11.733), Inches(1.3))
    achain.fill.solid()
    achain.fill.fore_color.rgb = RGBColor(241, 245, 249)
    achain.line.color.rgb = COLOR_BORDER_TEAL
    achain.line.width = Pt(1)

    tf_ac = achain.text_frame
    tf_ac.word_wrap = True
    tf_ac.vertical_anchor = MSO_ANCHOR.TOP
    tf_ac.margin_left = Inches(0.25)
    tf_ac.margin_top = Inches(0.18)
    tf_ac.margin_right = Inches(0.25)
    tf_ac.margin_bottom = Inches(0.18)

    p_ach = tf_ac.paragraphs[0]
    p_ach.text = "COMPLETE SYSTEM DATA FLOW:"
    p_ach.font.size = Pt(10)
    p_ach.font.bold = True
    p_ach.font.color.rgb = COLOR_TEAL_DARK
    p_ach.font.name = FONT_HEADING
    p_ach.alignment = PP_ALIGN.LEFT

    p_ac_flow = tf_ac.add_paragraph()
    p_ac_flow.text = "Patient / Staff Browser  ──▶  React Frontend  ──▶  FastAPI Endpoints  ──▶  AI Safety Pipeline  ──▶  SQLite / PostgreSQL  ──▶  Immutable Audit Trail"
    p_ac_flow.font.size = Pt(9.5)
    p_ac_flow.font.bold = True
    p_ac_flow.font.color.rgb = COLOR_TEXT_PRIMARY_DARK
    p_ac_flow.font.name = FONT_BODY
    p_ac_flow.alignment = PP_ALIGN.LEFT
    p_ac_flow.space_before = Pt(4)

    # =========================================================================
    # SLIDE 8 — SAFETY, PRIVACY & HUMAN OVERSIGHT
    # =========================================================================
    s8 = prs.slides.add_slide(blank_layout)
    add_slide_bg(s8, is_dark=False)
    add_header(s8, "Safety by Design: Human Oversight & Provenance", "SAFETY & GOVERNANCE", 8)

    pillars = [
        ("1. Clinician Supervision & Override", "Physicians retain complete clinical authority to confirm, modify, or override AI triage suggestions with mandatory clinical rationale notes."),
        ("2. Deterministic Red-Flag Rules", "Hardcoded medical safety rules immediately intercept acute cardiorespiratory and critical red flags, bypassing generic conversational delays."),
        ("3. Role-Based Access & Auditability", "Strict JWT-based RBAC partitions Patient self-intake from Nurse and Doctor workflows. All events are recorded in an immutable audit trail.")
    ]
    y_pil = 1.35
    for pil_t, pil_d in pillars:
        pil_card = s8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(y_pil), Inches(5.7), Inches(1.3))
        pil_card.fill.solid()
        pil_card.fill.fore_color.rgb = COLOR_WHITE
        pil_card.line.color.rgb = COLOR_BORDER_TEAL if "Override" in pil_t else COLOR_BORDER_LIGHT
        pil_card.line.width = Pt(1.2)

        tf = pil_card.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.TOP
        tf.margin_left = Inches(0.25)
        tf.margin_top = Inches(0.18)
        tf.margin_right = Inches(0.25)
        tf.margin_bottom = Inches(0.18)

        p_t = tf.paragraphs[0]
        p_t.text = pil_t
        p_t.font.size = Pt(11)
        p_t.font.bold = True
        p_t.font.color.rgb = COLOR_TEXT_PRIMARY_DARK
        p_t.font.name = FONT_HEADING
        p_t.alignment = PP_ALIGN.LEFT

        p_d = tf.add_paragraph()
        p_d.text = pil_d
        p_d.font.size = Pt(9.5)
        p_d.font.color.rgb = COLOR_TEXT_SECONDARY
        p_d.font.name = FONT_BODY
        p_d.alignment = PP_ALIGN.LEFT
        p_d.space_before = Pt(3)
        y_pil += 1.4

    gov_card = s8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(5.6), Inches(5.7), Inches(1.5))
    gov_card.fill.solid()
    gov_card.fill.fore_color.rgb = RGBColor(204, 251, 241)
    gov_card.line.color.rgb = COLOR_BORDER_TEAL
    gov_card.line.width = Pt(1.2)

    tf_gc = gov_card.text_frame
    tf_gc.word_wrap = True
    tf_gc.vertical_anchor = MSO_ANCHOR.TOP
    tf_gc.margin_left = Inches(0.25)
    tf_gc.margin_top = Inches(0.2)
    tf_gc.margin_right = Inches(0.25)
    tf_gc.margin_bottom = Inches(0.2)

    p_gh = tf_gc.paragraphs[0]
    p_gh.text = "CORE SAFETY PRINCIPLE:"
    p_gh.font.size = Pt(10)
    p_gh.font.bold = True
    p_gh.font.color.rgb = COLOR_TEAL_DARK
    p_gh.font.name = FONT_HEADING
    p_gh.alignment = PP_ALIGN.LEFT

    p_gd = tf_gc.add_paragraph()
    p_gd.text = "\"AI recommendations remain advisory; clinical decisions remain with qualified healthcare professionals.\""
    p_gd.font.size = Pt(10)
    p_gd.font.bold = True
    p_gd.font.color.rgb = COLOR_TEXT_PRIMARY_DARK
    p_gd.font.name = FONT_BODY
    p_gd.alignment = PP_ALIGN.LEFT
    p_gd.space_before = Pt(4)

    sc8a = os.path.join(SCREENSHOTS_DIR, "09_institutional_audit_trail.png")
    if os.path.exists(sc8a):
        s8.shapes.add_picture(sc8a, Inches(6.8), Inches(1.35), Inches(5.733), Inches(2.8))

    sc8b = os.path.join(SCREENSHOTS_DIR, "10_trust_center_governance.png")
    if os.path.exists(sc8b):
        s8.shapes.add_picture(sc8b, Inches(6.8), Inches(4.3), Inches(5.733), Inches(2.8))

    # =========================================================================
    # SLIDE 9 — VALIDATION, IMPACT & ROADMAP (BALANCED 3 COLUMNS)
    # =========================================================================
    s9 = prs.slides.add_slide(blank_layout)
    add_slide_bg(s9, is_dark=False)
    add_header(s9, "Validation, Expected Impact & Project Roadmap", "IMPACT & ROADMAP", 9)

    # Col 1: Expected Impact
    s1_card = s9.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.35), Inches(3.7), Inches(5.75))
    s1_card.fill.solid()
    s1_card.fill.fore_color.rgb = COLOR_WHITE
    s1_card.line.color.rgb = COLOR_BORDER_TEAL
    s1_card.line.width = Pt(1.5)

    tf1 = s1_card.text_frame
    tf1.word_wrap = True
    tf1.vertical_anchor = MSO_ANCHOR.TOP
    tf1.margin_left = Inches(0.25)
    tf1.margin_top = Inches(0.25)
    tf1.margin_right = Inches(0.25)
    tf1.margin_bottom = Inches(0.25)

    p1_h = tf1.paragraphs[0]
    p1_h.text = "SECTION 1 — EXPECTED IMPACT"
    p1_h.font.size = Pt(11)
    p1_h.font.bold = True
    p1_h.font.color.rgb = COLOR_TEAL_DARK
    p1_h.font.name = FONT_HEADING
    p1_h.alignment = PP_ALIGN.LEFT
    p1_h.space_after = Pt(10)

    impacts = [
        ("Faster Patient Structuring:", "Rapidly organizes unstructured complaints into clinical timelines."),
        ("Early Red-Flag Detection:", "Prioritizes cardiorespiratory emergencies before triage bottlenecks."),
        ("Multimodal Flexibility:", "Accommodates text, voice, and medical lab reports."),
        ("Seamless Clinical Handoff:", "Smooth data flow across Patient, Nurse, and Doctor."),
        ("Complete Traceability:", "Full audit history with clinician sign-off.")
    ]
    for lbl_i, desc_i in impacts:
        p = tf1.add_paragraph()
        p.text = f"✔ {lbl_i} "
        p.font.size = Pt(9.5)
        p.font.bold = True
        p.font.color.rgb = COLOR_TEXT_PRIMARY_DARK
        p.font.name = FONT_HEADING
        p.alignment = PP_ALIGN.LEFT
        
        r_i = p.add_run()
        r_i.text = desc_i
        r_i.font.size = Pt(9)
        r_i.font.bold = False
        r_i.font.color.rgb = COLOR_TEXT_SECONDARY
        r_i.font.name = FONT_BODY
        p.space_before = Pt(6)

    # Col 2: Real Project Validation Metrics
    s2_card = s9.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(4.8), Inches(1.35), Inches(3.7), Inches(5.75))
    s2_card.fill.solid()
    s2_card.fill.fore_color.rgb = COLOR_WHITE
    s2_card.line.color.rgb = COLOR_MINT
    s2_card.line.width = Pt(1.5)

    tf2 = s2_card.text_frame
    tf2.word_wrap = True
    tf2.vertical_anchor = MSO_ANCHOR.TOP
    tf2.margin_left = Inches(0.25)
    tf2.margin_top = Inches(0.25)
    tf2.margin_right = Inches(0.25)
    tf2.margin_bottom = Inches(0.25)

    p2_h = tf2.paragraphs[0]
    p2_h.text = "SECTION 2 — REAL VALIDATION"
    p2_h.font.size = Pt(11)
    p2_h.font.bold = True
    p2_h.font.color.rgb = COLOR_MINT
    p2_h.font.name = FONT_HEADING
    p2_h.alignment = PP_ALIGN.LEFT
    p2_h.space_after = Pt(6)

    p2_sub = tf2.add_paragraph()
    p2_sub.text = "Reported Automated Test Results:"
    p2_sub.font.size = Pt(9.5)
    p2_sub.font.bold = True
    p2_sub.font.color.rgb = COLOR_TEXT_PRIMARY_DARK
    p2_sub.font.name = FONT_HEADING
    p2_sub.alignment = PP_ALIGN.LEFT
    p2_sub.space_after = Pt(8)

    stats = [
        ("32", "Health AI functional & safety scenarios"),
        ("192", "Backend-connected executions"),
        ("192 Passed", "0 Failures recorded in test runner"),
        ("13 / 13", "E2E verification phases passed"),
        ("3,448 B", "Binary ReportLab referral PDF generated")
    ]
    for num_s, desc_s in stats:
        p_n = tf2.add_paragraph()
        p_n.text = f"• {num_s}: "
        p_n.font.size = Pt(10)
        p_n.font.bold = True
        p_n.font.color.rgb = COLOR_TEAL_DARK
        p_n.font.name = FONT_HEADING
        p_n.alignment = PP_ALIGN.LEFT

        r_d = p_n.add_run()
        r_d.text = desc_s
        r_d.font.size = Pt(9)
        r_d.font.bold = False
        r_d.font.color.rgb = COLOR_TEXT_SECONDARY
        r_d.font.name = FONT_BODY
        p_n.space_before = Pt(6)

    # Col 3: Future Scope & Roadmap
    s3_card = s9.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.8), Inches(1.35), Inches(3.733), Inches(5.75))
    s3_card.fill.solid()
    s3_card.fill.fore_color.rgb = COLOR_WHITE
    s3_card.line.color.rgb = COLOR_PURPLE
    s3_card.line.width = Pt(1.5)

    tf3 = s3_card.text_frame
    tf3.word_wrap = True
    tf3.vertical_anchor = MSO_ANCHOR.TOP
    tf3.margin_left = Inches(0.25)
    tf3.margin_top = Inches(0.25)
    tf3.margin_right = Inches(0.25)
    tf3.margin_bottom = Inches(0.25)

    p3_h = tf3.paragraphs[0]
    p3_h.text = "SECTION 3 — FUTURE ROADMAP"
    p3_h.font.size = Pt(11)
    p3_h.font.bold = True
    p3_h.font.color.rgb = COLOR_PURPLE
    p3_h.font.name = FONT_HEADING
    p3_h.alignment = PP_ALIGN.LEFT
    p3_h.space_after = Pt(10)

    futures = [
        ("Wider Multilingual Coverage:", "Expanded regional Indian dialect models."),
        ("Hospital EHR / ABDM Integration:", "FHIR-compliant health records sync."),
        ("Multi-Center Clinical Validation:", "Hospital pilot trials across departments."),
        ("Offline Low-Connectivity Sync:", "Edge SQLite replication for rural PHCs."),
        ("Wearable Vitals Streaming:", "Continuous IoT telemetry ingestion.")
    ]
    for lbl_f, desc_f in futures:
        p = tf3.add_paragraph()
        p.text = f"➔ {lbl_f} "
        p.font.size = Pt(9.5)
        p.font.bold = True
        p.font.color.rgb = COLOR_TEXT_PRIMARY_DARK
        p.font.name = FONT_HEADING
        p.alignment = PP_ALIGN.LEFT

        r_f = p.add_run()
        r_f.text = desc_f
        r_f.font.size = Pt(9)
        r_f.font.bold = False
        r_f.font.color.rgb = COLOR_TEXT_SECONDARY
        r_f.font.name = FONT_BODY
        p.space_before = Pt(6)

    # =========================================================================
    # SLIDE 10 — THANK YOU (HERO DARK)
    # =========================================================================
    s10 = prs.slides.add_slide(blank_layout)
    add_slide_bg(s10, is_dark=True)

    top_bar10 = s10.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(0.7), Inches(11.733), Inches(0.05))
    top_bar10.fill.solid()
    top_bar10.fill.fore_color.rgb = COLOR_BORDER_TEAL
    top_bar10.line.fill.background()

    t10 = s10.shapes.add_textbox(Inches(0.8), Inches(1.1), Inches(11.733), Inches(1.6))
    tf10 = t10.text_frame
    tf10.word_wrap = True
    tf10.margin_left = Inches(0)
    tf10.margin_top = Inches(0)

    p_ty = tf10.paragraphs[0]
    p_ty.text = "THANK YOU"
    p_ty.font.size = Pt(40)
    p_ty.font.bold = True
    p_ty.font.color.rgb = COLOR_TEXT_WHITE
    p_ty.font.name = FONT_HEADING
    p_ty.alignment = PP_ALIGN.LEFT

    p_q = tf10.add_paragraph()
    p_q.text = "Questions & Demonstration"
    p_q.font.size = Pt(17)
    p_q.font.color.rgb = COLOR_CYAN_LIGHT
    p_q.font.name = FONT_HEADING
    p_q.alignment = PP_ALIGN.LEFT
    p_q.space_before = Pt(2)

    card10_l = s10.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(2.95), Inches(6.0), Inches(3.7))
    card10_l.fill.solid()
    card10_l.fill.fore_color.rgb = COLOR_NAVY_SURFACE
    card10_l.line.color.rgb = COLOR_NAVY_BORDER
    card10_l.line.width = Pt(1)

    tf_f10 = card10_l.text_frame
    tf_f10.word_wrap = True
    tf_f10.vertical_anchor = MSO_ANCHOR.TOP
    tf_f10.margin_left = Inches(0.3)
    tf_f10.margin_top = Inches(0.3)
    tf_f10.margin_right = Inches(0.3)
    tf_f10.margin_bottom = Inches(0.3)

    p_proj = tf_f10.paragraphs[0]
    p_proj.text = "PROJECT:  SWASTHYA TRIAGE"
    p_proj.font.size = Pt(13)
    p_proj.font.bold = True
    p_proj.font.color.rgb = COLOR_CYAN_LIGHT
    p_proj.font.name = FONT_HEADING
    p_proj.alignment = PP_ALIGN.LEFT
    p_proj.space_after = Pt(10)

    ty_fields = [
        ("Team Name:", "[ADD]"),
        ("Problem Statement ID:", "[ADD]"),
        ("Live Website / QR Code:", "[ADD LATER]"),
        ("Technology Stack:", "FastAPI + Python + React 19 + SQLite")
    ]
    for idx, (lbl, val) in enumerate(ty_fields):
        p = tf_f10.add_paragraph()
        p.text = f"{lbl} "
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = COLOR_CYAN_LIGHT if idx < 3 else COLOR_MINT
        p.font.name = FONT_HEADING
        p.alignment = PP_ALIGN.LEFT

        r = p.add_run()
        r.text = val
        r.font.size = Pt(11)
        r.font.bold = False
        r.font.color.rgb = RGBColor(253, 224, 71) if idx < 3 else COLOR_TEXT_WHITE
        r.font.name = FONT_HEADING
        p.space_after = Pt(8)

    card10_r = s10.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.1), Inches(2.95), Inches(5.433), Inches(3.7))
    card10_r.fill.solid()
    card10_r.fill.fore_color.rgb = COLOR_NAVY_SURFACE
    card10_r.line.color.rgb = COLOR_BORDER_TEAL
    card10_r.line.width = Pt(1.5)

    tf_rc = card10_r.text_frame
    tf_rc.word_wrap = True
    tf_rc.vertical_anchor = MSO_ANCHOR.TOP
    tf_rc.margin_left = Inches(0.3)
    tf_rc.margin_top = Inches(0.3)
    tf_rc.margin_right = Inches(0.3)
    tf_rc.margin_bottom = Inches(0.3)

    p_rch = tf_rc.paragraphs[0]
    p_rch.text = "Swasthya Triage Highlights:"
    p_rch.font.size = Pt(13)
    p_rch.font.bold = True
    p_rch.font.color.rgb = COLOR_CYAN_LIGHT
    p_rch.font.name = FONT_HEADING
    p_rch.alignment = PP_ALIGN.LEFT
    p_rch.space_after = Pt(8)

    closing_points = [
        "100% Real Working Full-Stack Healthcare System",
        "Multimodal Intake (Text, Voice, Medical Reports OCR)",
        "Strictly Non-Diagnostic Clinician-Supervised AI",
        "Dedicated Workflows for Patient, Nurse, and Doctor",
        "192/192 Backend Automated Safety Tests Passed"
    ]
    for pt in closing_points:
        p = tf_rc.add_paragraph()
        p.text = f"★  {pt}"
        p.font.size = Pt(10)
        p.font.color.rgb = COLOR_TEXT_WHITE
        p.font.name = FONT_HEADING
        p.alignment = PP_ALIGN.LEFT
        p.space_before = Pt(4)

    prs.save(OUTPUT_PPTX)
    print(f"Professional presentation generated successfully: {OUTPUT_PPTX}")

if __name__ == "__main__":
    create_deck()
